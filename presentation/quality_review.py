# -*- coding: utf-8 -*-
"""调用 cli-anything-wps 技能自带的质量审查模块（5 维度审查标准），
对生成的 deck 进行 visual 维度打分（字号层级/内容密度/视觉占比）。
"""
import os
import sys

HERE = os.path.dirname(__file__)
SKILL_STYLES = os.path.join(HERE, "..", "skills", "cli-anything-wps", "styles")
sys.path.insert(0, SKILL_STYLES)

import design_presets               # noqa: E402
import quality_checks as qc         # noqa: E402
from pptx import Presentation       # noqa: E402

DECK = os.path.join(HERE, "废旧轮胎的回收与资源化利用.pptx")
PRESET = design_presets.PRESETS["academic"]


def role_of(size_pt):
    if size_pt >= 28:
        return "title"
    if size_pt >= 20:
        return "body"
    return "caption"


def extract_slide_elements(slide):
    elements = []
    for shp in slide.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip():
            sizes = [r.font.size.pt for p in shp.text_frame.paragraphs
                     for r in p.runs if r.font.size]
            fs = max(sizes) if sizes else 18
            elements.append({"type": "text", "fs": fs, "role": role_of(fs)})
        elif shp.shape_type is not None:
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            if shp.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.AUTO_SHAPE):
                elements.append({"type": "image" if shp.shape_type == MSO_SHAPE_TYPE.PICTURE
                                  else "box"})
    return elements


def main():
    prs = Presentation(DECK)
    all_slides = [extract_slide_elements(s) for s in prs.slides]
    report = qc.review_deck(all_slides, PRESET)
    print(f"预设: {PRESET.name}  审查页数: {len(all_slides)}")
    print(report["summary"])
    print(f"通过阈值(visual>=70): {'PASS' if report['overall_score'] >= 70 else 'FAIL'}")
    worst = sorted(report["per_slide"], key=lambda r: r["score"])[:5]
    if worst and worst[0]["score"] < 100:
        print("\n得分最低的页面：")
        for r in worst:
            if r["score"] < 100:
                print(f"  第{r['slide_index']+1:02d}页  {r['score']}分  "
                      f"{'; '.join(r['warnings']) if r['warnings'] else ''}")


if __name__ == "__main__":
    main()
